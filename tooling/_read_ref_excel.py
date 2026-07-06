"""Read the reference Excel and dump key sheets."""
import openpyxl, sys

wb = openpyxl.load_workbook(
    r'G:\Downloads\Household_Adaptive_TDEE_Planner_v7_Excel_Safe.xlsx',
    data_only=True,
)

sheets_to_read = [
    'Start Here', 'Weekly Dashboard', 'Refeed & Diet Break',
    'Dynamic Projection', 'Body Measurements', 'Daily Carryover Plan',
    'Exercise Log', 'Weight Tracker', 'Male Calculator',
]

for sn in sheets_to_read:
    ws = wb[sn]
    print(f'\n{"="*60}')
    print(f'SHEET: {sn}  (rows={ws.max_row}, cols={ws.max_column})')
    print(f'{"="*60}')
    for r in range(1, min(ws.max_row + 1, 55)):
        vals = []
        for c in range(1, min(ws.max_column + 1, 14)):
            v = ws.cell(row=r, column=c).value
            if v is not None:
                s = str(v)[:120].replace('\n', '\\n')
                vals.append(f'C{c}={s}')
        if vals:
            print(f'  R{r}: ' + ' | '.join(vals))
    if ws.max_row > 55:
        print(f'  ... ({ws.max_row - 55} more rows)')
